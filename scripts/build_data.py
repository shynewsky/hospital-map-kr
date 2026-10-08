"""Build static hospital data for Hospital Map KR from official HIRA files.

Usage:
  python scripts/build_data.py \
    --hira-zip path/to/hira-hospitals.zip \
    --equipment-csv path/to/hira-equipment.csv \
    --out data

The script uses only the Python standard library and openpyxl.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import shutil
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path

import openpyxl

ALLOWED_TYPES = {
    "상급종합",
    "종합병원",
    "병원",
    "정신병원",
    "요양병원",
    "의원",
    "보건의료원",
    "보건소",
    "보건지소",
    "보건진료소",
}

REGION_NAMES = {
    "110000": "서울",
    "210000": "부산",
    "220000": "인천",
    "230000": "대구",
    "240000": "광주",
    "250000": "대전",
    "260000": "울산",
    "290000": "세종",
    "310000": "경기",
    "320000": "강원",
    "330000": "충북",
    "340000": "충남",
    "350000": "전북",
    "360000": "전남",
    "370000": "경북",
    "380000": "경남",
    "390000": "제주",
}

EQUIPMENT_CODE_TO_KEY = {
    "B101": "xray",
    "B108": "ct",
    "B301": "mri",
}

TYPE_ORDER = {
    "상급종합": 1,
    "종합병원": 2,
    "병원": 3,
    "정신병원": 4,
    "요양병원": 5,
    "의원": 6,
    "보건의료원": 7,
    "보건소": 8,
    "보건지소": 9,
    "보건진료소": 10,
}


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def integer(value) -> int:
    if value in (None, ""):
        return 0
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def number(value):
    if value in (None, ""):
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    return parsed if math.isfinite(parsed) else None


def code_string(value, width=0) -> str:
    if value in (None, ""):
        return ""
    try:
        result = str(int(float(value)))
    except (TypeError, ValueError):
        result = str(value).strip()
    return result.zfill(width) if width else result


def public_id(ykiho: str) -> str:
    return hashlib.sha256(ykiho.encode("utf-8")).hexdigest()[:20]


def normalize_phone(value: str) -> str:
    phone = text(value)
    if not phone:
        return ""
    return phone


def find_file(root: Path, starts_with: str) -> Path:
    matches = [path for path in root.rglob("*.xlsx") if path.name.startswith(starts_with)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one XLSX starting with {starts_with!r}, found {len(matches)}")
    return matches[0]


def load_hospitals(basic_path: Path):
    workbook = openpyxl.load_workbook(basic_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows)]
    positions = {name: index for index, name in enumerate(headers)}

    required = [
        "암호화요양기호",
        "요양기관명",
        "종별코드",
        "종별코드명",
        "시도코드",
        "시도코드명",
        "시군구코드",
        "시군구코드명",
        "주소",
        "전화번호",
        "병원홈페이지",
        "총의사수",
        "의과일반의 인원수",
        "의과전문의 인원수",
        "좌표(X)",
        "좌표(Y)",
    ]
    missing = [name for name in required if name not in positions]
    if missing:
        raise RuntimeError(f"Missing basic columns: {missing}")

    hospitals = {}
    skipped_type = 0
    skipped_coordinate = 0

    for row in rows:
        institution_type = text(row[positions["종별코드명"]])
        if institution_type not in ALLOWED_TYPES:
            skipped_type += 1
            continue

        longitude = number(row[positions["좌표(X)"]])
        latitude = number(row[positions["좌표(Y)"]])
        if longitude is None or latitude is None or not (124 <= longitude <= 132 and 33 <= latitude <= 39.5):
            skipped_coordinate += 1
            continue

        ykiho = text(row[positions["암호화요양기호"]])
        if not ykiho:
            continue

        region_code = code_string(row[positions["시도코드"]], 6)
        region_name = REGION_NAMES.get(region_code, text(row[positions["시도코드명"]]))
        hospital = {
            "id": public_id(ykiho),
            "name": text(row[positions["요양기관명"]]),
            "typeCode": code_string(row[positions["종별코드"]]),
            "type": institution_type,
            "regionCode": region_code,
            "region": region_name,
            "districtCode": code_string(row[positions["시군구코드"]], 6),
            "district": text(row[positions["시군구코드명"]]),
            "address": text(row[positions["주소"]]),
            "phone": normalize_phone(row[positions["전화번호"]]),
            "website": text(row[positions["병원홈페이지"]]),
            "lat": round(latitude, 7),
            "lng": round(longitude, 7),
            "doctorTotal": integer(row[positions["총의사수"]]),
            "generalDoctorCount": integer(row[positions["의과일반의 인원수"]]),
            "specialistTotal": integer(row[positions["의과전문의 인원수"]]),
            "departments": [],
            "equipment": {"xray": 0, "ct": 0, "mri": 0},
            "designations": [],
            "_ykiho": ykiho,
        }
        hospitals[ykiho] = hospital

    workbook.close()
    return hospitals, {"skippedType": skipped_type, "skippedCoordinate": skipped_coordinate}


def attach_departments(hospitals, department_path: Path):
    workbook = openpyxl.load_workbook(department_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows)]
    positions = {name: index for index, name in enumerate(headers)}
    required = ["암호화요양기호", "진료과목코드", "진료과목코드명", "과목별 전문의수"]
    missing = [name for name in required if name not in positions]
    if missing:
        raise RuntimeError(f"Missing department columns: {missing}")

    department_catalog = {}
    attached = 0
    for row in rows:
        ykiho = text(row[positions["암호화요양기호"]])
        hospital = hospitals.get(ykiho)
        if hospital is None:
            continue
        department_code = code_string(row[positions["진료과목코드"]], 2)
        department_name = text(row[positions["진료과목코드명"]])
        if not department_code or not department_name:
            continue
        specialist_count = integer(row[positions["과목별 전문의수"]])
        hospital["departments"].append(
            {"code": department_code, "name": department_name, "specialists": specialist_count}
        )
        department_catalog[department_code] = department_name
        attached += 1

    workbook.close()
    for hospital in hospitals.values():
        hospital["departments"].sort(key=lambda item: (-item["specialists"], item["name"]))
    return department_catalog, attached


def attach_designations(hospitals, designation_path: Path):
    workbook = openpyxl.load_workbook(designation_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows)]
    positions = {name: index for index, name in enumerate(headers)}
    attached = 0
    for row in rows:
        ykiho = text(row[positions["암호화요양기호"]])
        hospital = hospitals.get(ykiho)
        if hospital is None:
            continue
        label = text(row[positions["검색코드명"]])
        if label and label not in hospital["designations"]:
            hospital["designations"].append(label)
            attached += 1
    workbook.close()
    return attached


def attach_equipment(hospitals, equipment_csv: Path):
    attached = 0
    matched_hospitals = set()
    with equipment_csv.open("r", encoding="cp949", errors="replace", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"암호화된 요양기호", "장비대분류코드", "장비수"}
        if not required.issubset(set(reader.fieldnames or [])):
            raise RuntimeError(f"Missing equipment columns. Found: {reader.fieldnames}")
        for row in reader:
            key = EQUIPMENT_CODE_TO_KEY.get(text(row.get("장비대분류코드")))
            if key is None:
                continue
            ykiho = text(row.get("암호화된 요양기호"))
            hospital = hospitals.get(ykiho)
            if hospital is None:
                continue
            quantity = integer(row.get("장비수"))
            if quantity <= 0:
                continue
            hospital["equipment"][key] += quantity
            attached += 1
            matched_hospitals.add(ykiho)
    return attached, len(matched_hospitals)


def export_data(hospitals, department_catalog, output_dir: Path):
    output_dir.mkdir(parents=True, exist_ok=True)
    regions_dir = output_dir / "regions"
    regions_dir.mkdir(parents=True, exist_ok=True)
    for existing in regions_dir.glob("*.json"):
        existing.unlink()

    by_region = defaultdict(list)
    bounds = {}
    for hospital in hospitals.values():
        ykiho = hospital.pop("_ykiho", None)
        del ykiho
        by_region[hospital["regionCode"]].append(hospital)

    region_metadata = []
    total_exported = 0
    for region_code, items in sorted(by_region.items()):
        items.sort(key=lambda item: (TYPE_ORDER.get(item["type"], 99), item["name"]))
        latitudes = [item["lat"] for item in items]
        longitudes = [item["lng"] for item in items]
        region_name = REGION_NAMES.get(region_code, items[0]["region"] if items else region_code)
        region_bounds = {
            "south": min(latitudes),
            "west": min(longitudes),
            "north": max(latitudes),
            "east": max(longitudes),
        }
        bounds[region_code] = region_bounds
        payload = {
            "region": {"code": region_code, "name": region_name},
            "sourceReferenceDate": "2026-06-30",
            "equipmentReferenceDate": "2024-12-31",
            "hospitals": items,
        }
        target = regions_dir / f"{region_code}.json"
        target.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        region_metadata.append(
            {
                "code": region_code,
                "name": region_name,
                "count": len(items),
                "url": f"data/regions/{region_code}.json",
                "bounds": region_bounds,
            }
        )
        total_exported += len(items)

    departments = [
        {"code": code, "name": name}
        for code, name in sorted(department_catalog.items(), key=lambda item: item[1])
    ]
    (output_dir / "departments.json").write_text(
        json.dumps({"departments": departments}, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )

    manifest = {
        "schemaVersion": "1.0.0",
        "datasetVersion": "2026-06",
        "generatedFrom": {
            "hospitalData": "HIRA 전국 병의원 및 약국 현황 2026.6",
            "equipmentData": "HIRA 의료장비 상세 현황 2024.12",
        },
        "totalHospitals": total_exported,
        "regions": region_metadata,
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return manifest


def main():
    parser = argparse.ArgumentParser()
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--hira-zip", type=Path)
    source.add_argument("--hira-dir", type=Path)
    parser.add_argument("--equipment-csv", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()

    def build_from(extraction_root: Path):
        basic_path = find_file(extraction_root, "1.")
        department_path = find_file(extraction_root, "5.")
        designation_path = find_file(extraction_root, "11.")
        hospitals, skipped = load_hospitals(basic_path)
        department_catalog, department_rows = attach_departments(hospitals, department_path)
        designation_rows = attach_designations(hospitals, designation_path)
        equipment_rows, equipment_hospitals = attach_equipment(hospitals, args.equipment_csv)
        manifest = export_data(hospitals, department_catalog, args.out)
        return skipped, department_rows, designation_rows, equipment_rows, equipment_hospitals, manifest

    if args.hira_dir:
        result = build_from(args.hira_dir)
    else:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="hospital-map-hira-", dir=args.out.parent) as temp_dir:
            extraction_root = Path(temp_dir)
            with zipfile.ZipFile(args.hira_zip) as archive:
                archive.extractall(extraction_root)
            result = build_from(extraction_root)

    skipped, department_rows, designation_rows, equipment_rows, equipment_hospitals, manifest = result

    print(
        json.dumps(
            {
                "exportedHospitals": manifest["totalHospitals"],
                "regions": len(manifest["regions"]),
                "departmentRows": department_rows,
                "designationRows": designation_rows,
                "equipmentRows": equipment_rows,
                "equipmentHospitals": equipment_hospitals,
                **skipped,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
